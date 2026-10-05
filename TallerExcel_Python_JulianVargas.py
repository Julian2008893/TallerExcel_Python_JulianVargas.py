import openpyxl
from tabulate import tabulate

# Cargar el archivo Excel
excel_dataframe = openpyxl.load_workbook("personas.xlsx")
dataframe = excel_dataframe.active

# Guardar los datos en una lista
data = []

for row in range(1, dataframe.max_row):
    _row = [row]
    for col in dataframe.iter_cols(1, dataframe.max_column):
        _row.append(col[row].value)
    data.append(_row)

headers = ["#", "Id", "Name", "Company", "Email", "MAC Address"]
headers_align = ("center",) * 6


# - RETOS PRÁCTICOS -

# Reto 1: Mostrar únicamente Name y Email
def reto1():
    print("\n--- RETO 1: Solo Name y Email ---")
    datos_reto1 = []
    for fila in data:
        datos_reto1.append([fila[2], fila[4]])
    
    print(tabulate(datos_reto1, headers=["Name", "Email"], tablefmt="fancy_grid"))

# Reto 2: Mostrar cuántos registros tiene el archivo
def reto2():
    print("\n--- RETO 2: Total de registros ---")
    total = len(data)
    print(f"El archivo tiene un total de {total} registros.")

# Reto 3: Filtrar personas cuya empresa contenga la palabra Group
def reto3():
    print("\n--- RETO 3: Filtrar empresas con 'Group' ---")
    filtro_group = []
    for fila in data:
        # Verificamos si la empresa tiene la palabra Group
        if fila[3] and "Group" in str(fila[3]):
            filtro_group.append(fila)
            
    print(tabulate(filtro_group, headers=headers, tablefmt="fancy_grid", colalign=headers_align))

# Reto 4: Pedir un ID con input() y mostrar solo esa persona
def reto4():
    print("\n--- RETO 4: Buscar persona por ID ---")
    id_buscar = input("Ingrese el ID de la persona: ")
    
    encontrado = []
    for fila in data:
        if str(fila[1]) == id_buscar:
            encontrado.append(fila)
            
    if len(encontrado) > 0:
        print(tabulate(encontrado, headers=headers, tablefmt="fancy_grid", colalign=headers_align))
    else:
        print("No se encontró a nadie con ese ID.")

# Reto 5: Crear un menú interactivo
def menu():
    opcion = ""
    while opcion != "6":
        print("\n" + "="*30)
        print("      MENÚ DE OPCIONES")
        print("="*30)
        print("1. Ver todos los datos")
        print("2. Ver Name y Email (Reto 1)")
        print("3. Contar registros (Reto 2)")
        print("4. Filtrar empresas 'Group' (Reto 3)")
        print("5. Buscar por ID (Reto 4)")
        print("6. Salir")
        
        opcion = input("Elija una opción (1-6): ")
        
        if opcion == "1":
            print(tabulate(data, headers=headers, tablefmt="fancy_grid", colalign=headers_align))
        elif opcion == "2":
            reto1()
        elif opcion == "3":
            reto2()
        elif opcion == "4":
            reto3()
        elif opcion == "5":
            reto4()
        elif opcion == "6":
            print("Saliendo del programa...")
        else:
            print("Opción incorrecta, intente de nuevo.")

# Ejecutar el menú principal
menu()